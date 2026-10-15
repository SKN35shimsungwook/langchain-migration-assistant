"""발표 PPT 만들기: python build_deck.py <출력.pptx> <화면 캡처.png>"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT, SHOT = sys.argv[1], sys.argv[2]

NAVY, NAVY2, LIGHT, WHITE = "14202E", "1E2D3F", "F6F5F1", "FFFFFF"
INK, GRAY, MUTED, LINE = "1B2433", "4A5363", "8A8F98", "E1E0D9"
BLUE, SKY, ORANGE, PEACH = "1F5FAE", "7FB2F0", "B4471B", "F2A27E"
FONT, MONO = "맑은 고딕", "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
page = [0]


def rgb(h):
    return RGBColor.from_string(h)


def new(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    page[0] += 1
    return s


def text(s, x, y, w, h, runs, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: 문자열, 또는 줄 목록 [ [(글자, {옵션}), ...], ... ]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [[(line, {})] for line in str(runs).split("\n")]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, opt in line:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = rgb(opt.get("color", color))
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", opt.get("font", font))
    return tb


def box(s, x, y, w, h, fill=WHITE, line=LINE, radius=True, dash=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.06
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1.25 if dash else 0.75)
        if dash:
            shp.line.dash_style = 4
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def arrow(s, x1, y1, x2, y2, color=MUTED, dash=False, width=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def line(s, x1, y1, x2, y2, color=MUTED, dash=False, width=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    return c


def head(s, kicker, title, dark=False, sub=None):
    text(s, 0.75, 0.62, 11.8, 0.3, kicker, size=13, bold=True, color=SKY if dark else BLUE)
    text(s, 0.75, 0.95, 11.8, 0.8, title, size=32, bold=True, color=LIGHT if dark else INK)
    if sub:
        text(s, 0.75, 1.72, 11.8, 0.5, sub, size=16, color="B9C3CF" if dark else GRAY)


def foot(s, dark=False, extra=""):
    text(s, 0.75, 6.92, 11.8, 0.3, f"LangChain 마이그레이션 도우미 · {page[0]}{extra}", size=11,
         color="8A9AAD" if dark else MUTED)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def table(s, x, y, w, rows, col_w, size=15, row_h=0.46, head_fill="ECEBE5", first_bold=False, colors=None):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(style)
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # No Style, Table Grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(r):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(head_fill if i == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.text = ""
            run = tf.paragraphs[0].add_run()
            run.text = val
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = i == 0 or (first_bold and j == 0)
            c = INK
            if colors and i > 0 and colors.get(j):
                c = colors[j]
            run.font.color.rgb = rgb(c)
            rPr = run._r.get_or_add_rPr()
            ea = rPr.makeelement(qn("a:ea"), {"typeface": FONT})
            rPr.append(ea)
    return tbl


def big_card(s, x, y, w, h, num, label, src=None, num_color=BLUE, num_size=40):
    box(s, x, y, w, h)
    text(s, x + 0.3, y + 0.25, w - 0.6, 0.8, num, size=num_size, bold=True, color=num_color)
    text(s, x + 0.3, y + 1.1, w - 0.6, 0.9, label, size=15, color=INK, spacing=1.3)
    if src:
        text(s, x + 0.3, y + h - 0.5, w - 0.6, 0.3, src, size=11, color=MUTED)


# 1. 표지 -------------------------------------------------------------------
s = new(NAVY)
text(s, 0.75, 0.75, 11.8, 0.4, "SKN35 · LLM 연동 내외부 문서 기반 질의응답 시스템", size=15, bold=True, color=SKY)
text(s, 0.75, 2.25, 11.8, 2.0, "LangChain 버전\n마이그레이션 도우미", size=54, bold=True, color="F2F1EC", spacing=1.1)
text(s, 0.75, 4.35, 11.8, 0.6, "옛 버전 코드를 원하는 버전에서 돌아가는 코드로, 근거와 검증까지 함께", size=21, color="B9C3CF")
text(s, 0.75, 6.55, 11.8, 0.4, "SKN35 3rd 4Team · 발표 [발표자 이름] · 2026-10-15", size=14, color="B9C3CF")
notes(s, "안녕하세요, SKN35 3기 4팀입니다. 저희는 옛 버전 LangChain 코드를 넣고 바꿀 버전을 고르면, "
         "그 버전에서 돌아가는 코드로 고쳐 주고 정말 돌아가는지 확인까지 해 주는 챗봇을 만들었습니다. "
         "핵심은 '근거'와 '검증' 두 가지입니다.")

# 2. 문제 -------------------------------------------------------------------
s = new(LIGHT)
head(s, "문제", "ChatGPT가 준 LangChain 코드, 실행하면 깨집니다")
box(s, 0.75, 2.0, 6.6, 3.1, fill=NAVY, line=None)
text(s, 1.05, 2.25, 6.0, 0.3, "# LLM이 흔히 주는 코드", size=13, color="8FA0B3", font=MONO)
text(s, 1.05, 2.65, 6.0, 0.8, "from langchain.chat_models import ChatOpenAI\nfrom langchain.chains import LLMChain",
     size=15, color="E8ECF1", font=MONO, spacing=1.4)
line(s, 1.05, 3.55, 7.05, 3.55, color="34445A", width=0.75)
text(s, 1.05, 3.75, 6.0, 0.3, "# LangChain 1.4.4에서 실행 결과", size=13, color="8FA0B3", font=MONO)
text(s, 1.05, 4.15, 6.0, 0.8, "ImportError: cannot import name 'ChatOpenAI'\nModuleNotFoundError: No module named 'langchain.chains'",
     size=13, color=PEACH, font=MONO, spacing=1.4)
text(s, 7.85, 2.0, 4.8, 0.5, "첫 줄부터 에러", size=22, bold=True)
text(s, 7.85, 2.5, 4.8, 1.0, "LangChain 1.0에서 패키지가 나뉘고, 옛 체인과 에이전트가 빠졌습니다.", size=17, color=GRAY, spacing=1.35)
text(s, 7.85, 3.7, 4.8, 0.5, "더 큰 문제", size=22, bold=True)
text(s, 7.85, 4.2, 4.8, 0.8, [[("LLM은 자기 답이 틀렸다는 것을 ", {}), ("모릅니다.", {"bold": True, "color": INK})]],
     size=17, color=GRAY, spacing=1.35)
foot(s)
notes(s, "LangChain을 LLM에게 물어보면 복잡한 질문일수록 옛 버전 코드를 줍니다. 왼쪽은 그 코드를 LangChain 1.4.4에서 "
         "실제로 실행한 결과입니다. 첫 줄부터 에러가 납니다. 더 큰 문제는 LLM이 이게 틀렸다는 걸 모른다는 점입니다.")

# 3. 원인 -------------------------------------------------------------------
s = new(LIGHT)
head(s, "원인: 버전 지체 환각", "인터넷에는 옛 버전 코드가 더 많습니다",
     sub="LLM은 데이터가 많은 쪽을 따라갑니다. 라이브러리는 빨리 바뀌어도, 학습 데이터 속 코드는 대부분 과거 버전입니다.")
for i, (n, lab, src) in enumerate([
        ("42.55%", "문서가 부족할 때, LLM 코드 중 목표 버전에서 실행되는 비율", "When LLMs Lag Behind, 2026"),
        ("66.36%", "구조화된 문서를 줘도 실행되는 비율", "When LLMs Lag Behind, 2026"),
        ("48~51%", "최신 상용 모델의 버전 지정 코드 생성 성공률", "GitChameleon 2.0, 2025")]):
    big_card(s, 0.75 + i * 4.0, 2.55, 3.75, 2.4, n, lab, src)
foot(s)
notes(s, "이 현상을 '버전 지체 환각'이라고 부르겠습니다. 문서가 부족하면 LLM 코드 중 약 43%만 실행되고, 문서를 줘도 66% 정도입니다. "
         "문서만 주는 것으로는 부족하고, 실제로 확인하는 단계가 필요하다는 뜻입니다. (출처: arXiv 2604.09515, arXiv 2507.12367)")

# 4. 해결 -------------------------------------------------------------------
s = new(LIGHT)
head(s, "해결", "근거로 고치고, 검사로 한 번 더 확인합니다")
table(s, 0.75, 2.0, 11.83, [
    ["구분", "일반 LLM", "우리 챗봇"],
    ["근거", "학습 데이터 (옛 버전이 다수)", "import · 메서드 대응표 + 최신 공식 문서 검색 (RAG)"],
    ["확인", "없음", "verifier가 8가지 검사 (문법, import 존재, 남은 옛 import 등)"],
    ["틀렸을 때", "모름", "실패 이유로 근거를 다시 찾고, 틀린 근거는 빼고 다시 고침 (최대 3회)"],
    ["답변", "코드만", "고친 코드 + 바뀐 줄마다 근거 ID + 시도 횟수"],
    ["버전", "대부분 최신으로만", "지금 버전 → 바꿀 버전 선택 (0.0.300 ~ 1.x, 되돌리기 포함)"],
], [1.9, 3.6, 6.33], size=15, row_h=0.62, first_bold=True, colors={1: GRAY})
foot(s)
notes(s, "저희 해결책은 이중 검증입니다. 첫째, 기억이 아니라 대응표와 최신 공식 문서를 근거로 고칩니다. 둘째, 고친 코드를 verifier가 "
         "8가지로 검사합니다. 틀리면 실패 이유를 들고 근거를 다시 찾되, 이미 틀린 근거는 빼고 다시 고칩니다. 최대 3번입니다. "
         "그래서 답변에는 코드뿐 아니라 바뀐 줄마다 근거 ID와 몇 번 만에 통과했는지가 함께 나옵니다.")

# 5. 데이터 -----------------------------------------------------------------
s = new(LIGHT)
head(s, "근거 데이터 · 2번", "고칠 때 볼 근거를 미리 만들어 두었습니다")
for i, (n, lab) in enumerate([("9,155줄", "import 대응표\n옛 경로 → 1.x 경로"),
                              ("42개", "메서드 대응표\n.run() → .invoke() 등"),
                              ("295개", "공식 문서\n최신 271 · 가이드 24"),
                              ("11,945개", "검색용 조각 (청크)\nChroma + BM25 검색")]):
    big_card(s, 0.75 + i * 3.0, 2.0, 2.8, 2.15, n, lab, num_size=32)
text(s, 0.75, 4.5, 11.8, 0.4, "출처 4가지", size=17, bold=True)
table(s, 0.75, 4.95, 11.83, [
    ["PyPI 패키지 코드", "langchain-cli 변환 규칙", "최신 공식 문서", "0.3.30 옛 마이그레이션 가이드"],
    ["버전별 도구 목록 23개", "공식 자동 변환 규칙", "임베딩: text-embedding-3-small", "옛 방식 → 새 방식 설명"],
], [2.95, 2.95, 2.95, 2.98], size=13, row_h=0.5)
foot(s, extra=" · 2026-10-14 팀 저장소 dev 기준")
notes(s, "2번이 만든 근거 데이터입니다. 옛 import 경로를 1.x 경로로 잇는 대응표가 9,155줄, .run 같은 옛 메서드를 잇는 메서드 대응표가 "
         "42개 있습니다. 최신 공식 문서 271개와 옛 마이그레이션 가이드 24개를 11,945개 조각으로 나눠 Chroma에 넣었고, "
         "단어 검색 BM25와 벡터 검색을 합친 하이브리드 검색으로 찾습니다. 버전별로 어떤 도구가 있는지도 23개 패키지 버전에서 뽑아 두었습니다.")

# 6. 시스템 구조 --------------------------------------------------------------
s = new(LIGHT)
head(s, "시스템 구조", "네 단계를 거치고, 실패하면 되돌아갑니다")
steps = [("① analyzer", "1번", "바꿀 부분 찾기"), ("② retriever", "2번", "대응표 정확 매칭\n+ 하이브리드 검색"),
         ("③ rewriter", "3번", "근거 ID를 붙여\n코드 다시 쓰기"), ("④ verifier", "3번", "LLM 없이\n8가지 검사"),
         ("결과 → 화면", "4번", "결과 패널 표시\nTiDB에 대화 저장")]
bw, gap, y0 = 2.05, 0.39, 2.35
for i, (name, who, desc) in enumerate(steps):
    x = 0.75 + i * (bw + gap)
    hl = i == 3
    box(s, x, y0, bw, 1.75, fill=BLUE if hl else WHITE, line=None if hl else LINE)
    text(s, x + 0.2, y0 + 0.18, bw - 0.4, 0.3, f"담당 {who}", size=12, color="D6E5F8" if hl else MUTED)
    text(s, x + 0.2, y0 + 0.48, bw - 0.4, 0.4, name, size=17, bold=True, color=WHITE if hl else INK, font=FONT)
    text(s, x + 0.2, y0 + 0.92, bw - 0.4, 0.8, desc, size=13, color="E4EEFB" if hl else GRAY, spacing=1.25)
    if i < 4:
        arrow(s, x + bw + 0.04, y0 + 0.87, x + bw + gap - 0.04, y0 + 0.87)
vx = 0.75 + 3 * (bw + gap) + bw / 2
rx = 0.75 + 1 * (bw + gap) + bw / 2
line(s, vx, y0 + 1.75, vx, 4.75, color=ORANGE, dash=True)
line(s, vx, 4.75, rx, 4.75, color=ORANGE, dash=True)
arrow(s, rx, 4.75, rx, y0 + 1.8, color=ORANGE, dash=True)
text(s, rx - 1.0, 4.85, vx - rx + 2.0, 0.4, "검증 실패 → 실패 이유 + 틀린 근거 제외로 다시 검색 (최대 3회)", size=14, bold=True,
     color=ORANGE, align=PP_ALIGN.CENTER)
text(s, 0.75, 5.75, 11.8, 0.4, "연결 · 재시도 판단: 1번 (LangGraph)   ·   변환 결과 기준: LangChain 1.4.4", size=14, color=GRAY)
foot(s)
notes(s, "질문이 들어오면 네 단계를 거칩니다. 1번이 만든 analyzer가 코드에서 바꿀 부분을 찾고, 2번의 retriever가 대응표와 문서에서 "
         "근거를 찾습니다. 3번의 rewriter가 바뀐 줄마다 근거 ID를 붙여 코드를 다시 쓰고, 핵심인 verifier가 LLM 없이 8가지를 검사합니다. "
         "실패하면 실패 이유를 들고 retriever로 돌아가는데, 이미 틀린 근거는 빼고 다시 찾습니다. 최대 3번입니다. "
         "단계 연결과 다시 갈지 판단은 1번이 LangGraph로 맡았고, 결과는 제가 만든 화면에 보여 주고 대화는 TiDB에 저장합니다.")

# 7. verifier ---------------------------------------------------------------
s = new(LIGHT)
head(s, "검증 · 3번", "verifier는 LLM 없이 8가지를 검사합니다")
checks = [("1", "문법", "파이썬 문법 통과"), ("2", "근거 ID", "보여 준 근거만 썼는지"),
          ("3", "변경 기록", "기록과 실제 코드가 같은지"), ("4", "남은 옛 import", "옛 경로가 남았는지"),
          ("5", "import 존재", "1.x에 그 경로가 진짜 있는지"), ("6", "남은 옛 메서드", ".run() 같은 옛 호출 (경고)"),
          ("7", "안 쓰는 import", "필요 없는 import가 남았나"), ("8", "반환값 보존", "결과 모양이 그대로인지")]
for i, (n, t, d) in enumerate(checks):
    col, row = i % 4, i // 4
    x, y = 0.75 + col * 3.0, 2.0 + row * 2.0
    box(s, x, y, 2.8, 1.75)
    text(s, x + 0.25, y + 0.2, 2.3, 0.4, n, size=22, bold=True, color=BLUE)
    text(s, x + 0.25, y + 0.68, 2.3, 0.4, t, size=17, bold=True)
    text(s, x + 0.25, y + 1.1, 2.3, 0.6, d, size=13, color=GRAY)
text(s, 0.75, 6.1, 11.8, 0.4, "오류 → 다시 시도할 이유로 돌려보냄   ·   경고 → 결과에 표시만 함", size=14, color=GRAY)
foot(s)
notes(s, "verifier가 이 프로젝트의 핵심입니다. LLM에게 맞았는지 묻지 않고, 코드로 8가지를 검사합니다. 문법, 근거 ID가 이번에 실제로 보여 준 "
         "근거인지, 변경 기록과 실제 코드가 같은지, 옛 import가 남았는지, 새 경로가 1.x에 진짜 있는지, 옛 메서드가 남았는지, 안 쓰는 import, "
         "반환값이 그대로인지입니다. 오류는 재시도 이유가 되고, 경고는 결과에 표시합니다. LLM에게 검증을 맡기지 않은 이유는 검증하는 LLM도 "
         "같은 옛 지식을 갖고 있기 때문입니다.")

# 8. 화면 -------------------------------------------------------------------
s = new(LIGHT)
head(s, "화면 · 4번", "결과와 근거를 한 화면에서 봅니다")
box(s, 0.75, 2.0, 7.9, 3.55, fill="ECEBE5", line=None)
s.shapes.add_picture(SHOT, Inches(0.85), Inches(2.1), width=Inches(7.7))
text(s, 0.75, 5.65, 7.9, 0.5, "화면 시험 캡처 · 변환은 그래프 연결 전 임시 결과, 근거 ID(imap-…)는 2번 대응표 실제 값",
     size=12, color=MUTED)
feats = [("대화형 + 세션 저장", "ChatGPT처럼 대화 목록, 새 대화마다 세션 하나 (TiDB)"),
         ("버전 고르기", "지금 코드 버전(자동 감지) → 바꿀 버전, 되돌리기 포함"),
         ("결과 패널", "검증 배지 · 전후 비교 · 바뀐 줄마다 근거 ID"),
         ("예시 질문 35개", "사이드바에서 골라 바로 보내기")]
for i, (t, d) in enumerate(feats):
    y = 2.0 + i * 0.98
    line(s, 8.95, y, 12.58, y, color="D4D3CC", width=1)
    text(s, 8.95, y + 0.12, 3.63, 0.4, t, size=16, bold=True, color=BLUE if i == 2 else INK)
    text(s, 8.95, y + 0.48, 3.63, 0.5, d, size=12, color=GRAY)
foot(s)
notes(s, "제가 맡은 화면입니다. ChatGPT처럼 왼쪽에 대화 목록이 있고, '새 대화'를 누를 때마다 세션이 하나 생겨서 그 안의 질문은 TiDB에 "
         "이어서 저장됩니다. 위에서 지금 코드 버전과 바꿀 버전을 고르는데, 최신으로 올리는 것뿐 아니라 옛 버전으로 되돌리는 것도 됩니다. "
         "오른쪽 결과 패널에는 검증 배지, 고치기 전과 후, 그리고 바뀐 줄마다 근거 ID가 나옵니다. 지금 캡처의 근거 ID는 2번 대응표의 실제 값이고, "
         "변환 결과는 그래프 연결 전 임시 결과입니다. [그래프 연결 후 캡처로 교체]")

# 9. 실험 설계 ---------------------------------------------------------------
s = new(LIGHT)
head(s, "실험 설계 · 4번", "같은 질문 35개, 같은 조건으로 비교했습니다")
cards = [("1.x로 올리기 29개", "10가지 유형: import 경로 · 로더 · 임베딩 · LLMChain → LCEL · 체인\n· RAG · 메모리 · 에이전트 · 출력 파서 · 실무 상황"),
         ("버전 오가기 6개", "1.x → 0.x 되돌리기, 0.3 → 0.1 · 0.0.354, 0.0.300 → 0.2 중간 단계"),
         ("난이도", "쉬움 7 · 중간 18 · 어려움 10\n35개 모두 실제로 깨지는지 미리 확인 (35/35)")]
for i, (t, d) in enumerate(cards):
    y = 2.0 + i * 1.45
    box(s, 0.75, y, 5.6, 1.3)
    text(s, 1.0, y + 0.15, 5.1, 0.4, t, size=17, bold=True)
    text(s, 1.0, y + 0.55, 5.1, 0.75, d, size=12, color=GRAY, spacing=1.25)
table(s, 6.7, 2.0, 5.88, [
    ["항목", "규칙"],
    ["비교 대상", "일반 LLM vs 우리 챗봇"],
    ["모델", "양쪽 같은 [모델명], temperature 0"],
    ["일반 LLM 질문", "같은 요청 문장 + 같은 코드"],
    ["채점", "verifier와 같은 검사 기준"],
    ["최대 시도", "3회 (우리 챗봇)"],
    ["지표", "검증 통과율 · 근거 포함률 · 재시도 효과"],
], [1.7, 4.18], size=14, row_h=0.56, first_bold=True)
foot(s)
notes(s, "공정하게 비교하려고 규칙을 먼저 정했습니다. 실제 사용자가 보낼 법한 요청 문장과 옛 코드로 질문 35개를 만들었습니다. "
         "1.x로 올리는 질문 29개는 10가지 유형이고, 옛 버전으로 되돌리거나 중간 버전으로만 올리는 질문 6개를 더했습니다. "
         "35개 모두 실제로 깨지는 코드인지, 기대 답이 실제로 import 되는지 미리 확인했습니다. 일반 LLM과 우리 챗봇은 같은 모델, 같은 질문을 쓰고, "
         "채점은 verifier와 같은 기준으로 했습니다. 지표는 검증 통과율, 바뀐 줄에 근거가 붙은 비율, 재시도로 오른 통과율 세 가지입니다.")

# 10. 데모 ------------------------------------------------------------------
s = new(NAVY)
head(s, "데모", "LLMChain 코드가 통과하기까지", dark=True)
demo = [("1", SKY, "옛 코드 입력", "0.0.300 LLMChain 코드,\n바꿀 버전 LangChain 1.x"),
        ("2", PEACH, "1 · 2회차 검증 실패", "실패 이유를 들고 근거를 다시 찾음\n(틀린 근거는 제외) [실제 결과로 수정]"),
        ("3", SKY, "3회차 통과", "prompt | llm | parser 형태로 변환\n바뀐 줄마다 근거 ID")]
for i, (n, c, t, d) in enumerate(demo):
    x = 0.75 + i * 4.0
    box(s, x, 2.0, 3.75, 2.6, fill=NAVY2, line=None)
    text(s, x + 0.3, 2.2, 3.1, 0.7, n, size=36, bold=True, color=c)
    text(s, x + 0.3, 3.0, 3.1, 0.45, t, size=19, bold=True, color="F2F1EC")
    text(s, x + 0.3, 3.5, 3.1, 1.0, d, size=13, color="B9C3CF", spacing=1.3)
text(s, 0.75, 5.0, 11.8, 0.5, "화면에서: 검증 배지(몇 번 만에 통과) · 전후 비교 · 변경표와 근거 ID · 대화 다시 열기", size=15, color="B9C3CF")
foot(s, dark=True, extra=" · 현장 데모 실패에 대비해 녹화본 준비")
notes(s, "이제 화면으로 직접 보여드리겠습니다. 0.0.300의 LLMChain 코드를 넣고 바꿀 버전으로 LangChain 1.x를 고르겠습니다. "
         "첫 번째와 두 번째 시도는 검증에서 실패하고, 실패 이유를 들고 근거를 다시 찾습니다. 세 번째에 prompt | llm | StrOutputParser() "
         "형태로 바뀌며 통과합니다. 결과 패널에서 몇 번 만에 통과했는지, 바뀐 줄마다 어떤 근거를 썼는지 볼 수 있습니다. [실제 실행 결과에 맞게 수정]")

# 11. 결과 ------------------------------------------------------------------
s = new(LIGHT)
head(s, "결과", "[결과 한 줄 요약을 여기에]")
text(s, 0.75, 2.2, 4.2, 0.4, "우리 챗봇 · 검증 통과율", size=15, bold=True, color=BLUE)
text(s, 0.75, 2.6, 4.2, 1.0, "[__]%", size=60, bold=True)
text(s, 0.75, 3.75, 4.2, 0.4, "일반 LLM · 검증 통과율", size=15, bold=True, color=ORANGE)
text(s, 0.75, 4.15, 4.2, 1.0, "[__]%", size=60, bold=True)
text(s, 0.75, 5.35, 4.2, 0.4, "근거 포함률 [__]%  ·  질문 35개 기준", size=15, color=GRAY)
box(s, 5.3, 2.0, 7.28, 4.4, line="C3C2B7", dash=True)
text(s, 5.3, 3.9, 7.28, 0.8, "[그래프 삽입]\n난이도 · 유형별 검증 통과율", size=16, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
foot(s)
notes(s, "결과입니다. 검증 통과율은 일반 LLM이 [__]%, 우리 챗봇이 [__]%였습니다. 특히 [__] 유형에서 차이가 가장 컸습니다. "
         "그리고 우리 챗봇이 바꾼 줄 중 [__]%에 근거 ID가 붙어 있었습니다. [실험 결과를 보고 제목과 설명 수정]")

# 12. 재시도 효과 -------------------------------------------------------------
s = new(LIGHT)
head(s, "검증 루프 효과", "다시 고치는 루프가 통과율을 [+__%p] 올렸습니다")
box(s, 0.75, 2.0, 7.4, 4.4, line="C3C2B7", dash=True)
text(s, 0.75, 3.9, 7.4, 0.8, "[그래프 삽입]\n1회차 → 2회차 → 3회차 누적 통과율", size=16, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
for i, (lab, val, c) in enumerate([("1회차 통과율", "[__]%", INK), ("최종 통과율 (최대 3회)", "[__]%", BLUE),
                                   ("통과까지 평균 시도", "[__]회", INK)]):
    y = 2.0 + i * 1.5
    line(s, 8.6, y, 12.58, y, color="D4D3CC", width=1)
    text(s, 8.6, y + 0.15, 4.0, 0.4, lab, size=15, color=GRAY)
    text(s, 8.6, y + 0.55, 4.0, 0.8, val, size=40, bold=True, color=c)
foot(s)
notes(s, "첫 시도만 보면 통과율이 [__]%인데, 실패 이유로 근거를 다시 찾아 고치면 최종 [__]%까지 올라갑니다. "
         "확인하고 다시 고치는 루프가 [__]%p를 만들었습니다. 일반 LLM에는 이 루프가 없어서 틀려도 그대로 끝납니다.")

# 13. 발견 ------------------------------------------------------------------
s = new(LIGHT)
head(s, "만들면서 발견", "통과했다고 끝이 아니었습니다")
finds = [("지원 종료 중인 패키지", "langchain-community는 import만 해도 '지원 종료 중' 경고를 냅니다. 지금 통과해도 곧 깨질 수 있어요.",
          "langchain-community 0.4.2 확인"),
         ("classic 우회", "LLMChain을 langchain_classic에서 가져오면 import는 통과하지만, 최신 방식(LCEL)은 아닙니다.",
          "대응표에 '옛 창고(classic)'로 따로 표시"),
         ("옛 버전끼리는 그대로", "0.0.300 ~ 0.3.30 사이에는 자주 쓰는 import 경로가 거의 같습니다. 크게 바뀌는 것은 1.x와의 사이입니다.",
          "버전별 도구 목록 비교로 확인")]
for i, (t, d, src) in enumerate(finds):
    x = 0.75 + i * 4.0
    box(s, x, 2.0, 3.75, 3.4)
    text(s, x + 0.3, 2.25, 3.15, 0.5, t, size=19, bold=True)
    text(s, x + 0.3, 2.85, 3.15, 1.9, d, size=14, color=GRAY, spacing=1.4)
    text(s, x + 0.3, 4.85, 3.15, 0.4, src, size=11, color=MUTED)
foot(s)
notes(s, "만들면서 예상하지 못한 것을 세 가지 발견했습니다. 첫째, langchain-community는 import만 해도 지원 종료 경고가 나옵니다. "
         "둘째, 옛 체인을 langchain_classic에서 가져오면 import는 되지만 최신 방식은 아닙니다. 그래서 대응표와 화면에서 '옛 창고'로 따로 표시했습니다. "
         "셋째, 버전 되돌리기를 만들면서 버전별 도구 목록을 비교해 보니, 0.0.300부터 0.3.30까지는 자주 쓰는 경로가 거의 같았습니다. "
         "크게 바뀌는 것은 1.x와의 사이입니다.")

# 14. 한계 ------------------------------------------------------------------
s = new(LIGHT)
head(s, "한계와 다음 단계", "통과가 곧 정답은 아닙니다")
table(s, 0.75, 2.0, 11.83, [
    ["한계", "다음 단계"],
    ["import · 메서드까지 검사하지만, 실행 결과가 맞는지는 확인 못 함", "가짜 LLM을 넣어 코드 전체를 실제로 실행"],
    ["다시 쓰기(rewriter)는 1.x로 올리기 중심, 되돌리기는 화면 쪽 도구 목록 비교", "목표 버전을 그래프 전체에 넘겨 모든 방향 지원"],
    ["질문 35개, 팀이 직접 구축", "공개 저장소의 실사용 코드로 확장"],
    ["LangChain만 대상", "pydantic, OpenAI SDK 등 빠르게 바뀌는 라이브러리로 확장"],
], [6.3, 5.53], size=15, row_h=0.7)
foot(s)
notes(s, "한계도 말씀드리겠습니다. 가장 큰 한계는 통과가 곧 정답은 아니라는 점입니다. import와 메서드까지는 검사하지만, 함수에 넣는 값이나 "
         "실행 결과가 맞는지는 잡지 못합니다. 또 다시 쓰기는 1.x로 올리는 쪽이 중심이고, 옛 버전으로 되돌리는 것은 지금 화면 쪽에서 도구 목록을 "
         "비교해 처리합니다. 다음 단계로 가짜 LLM을 넣은 실제 실행 테스트, 목표 버전을 그래프 전체에 넘기기, 테스트셋과 대상 라이브러리 확장을 생각하고 있습니다.")

# 15. 가치 ------------------------------------------------------------------
s = new(BLUE)
text(s, 0.75, 0.62, 11.8, 0.3, "가치", size=13, bold=True, color="D6E5F8")
text(s, 0.75, 1.0, 11.8, 1.8, "라이브러리가 바뀔 때마다,\n근거와 함께 확인된 코드", size=42, bold=True, color=WHITE, spacing=1.1)
for i, (t, d) in enumerate([("시간 절감", "에러를 하나씩 검색하는 대신, 근거가 붙은 답을 바로 확인"),
                            ("신뢰", "바뀐 줄마다 근거 ID, 몇 번 만에 검사를 통과했는지 함께 제시"),
                            ("확장", "대응표 · 문서 · 검사 환경만 바꾸면 다른 라이브러리에도 적용")]):
    x = 0.75 + i * 4.0
    line(s, x, 4.7, x + 3.75, 4.7, color="8DB6EA", width=1.5)
    text(s, x, 4.9, 3.75, 0.45, t, size=19, bold=True, color=WHITE)
    text(s, x, 5.4, 3.75, 0.9, d, size=14, color="E4EEFB", spacing=1.35)
text(s, 0.75, 6.92, 11.8, 0.3, f"LangChain 마이그레이션 도우미 · {page[0]}", size=11, color="D6E5F8")
notes(s, "가치를 세 가지로 정리하겠습니다. 첫째, 개발자가 버전 에러를 하나씩 검색하는 시간을 줄여 줍니다. 둘째, 바뀐 줄마다 근거와 검사 결과가 "
         "붙어 있어서 믿고 쓸 수 있습니다. 셋째, 구조가 대응표와 문서, 검사 환경에만 의존하기 때문에 다른 라이브러리에도 그대로 적용할 수 있습니다.")

# 16. 팀 · 마무리 -------------------------------------------------------------
s = new(NAVY)
text(s, 0.75, 0.62, 11.8, 0.3, "팀", size=13, bold=True, color=SKY)
team = [("1번", "그래프 · 백엔드", "LangGraph 연결 · analyzer\n재시도 판단 · dev 병합"),
        ("2번", "데이터 · 검색", "대응표 · 청킹 · Chroma\n하이브리드 검색 · retriever"),
        ("3번", "변환 · 검증", "State · rewriter\nverifier 8가지 검사"),
        ("4번", "화면 · 평가 · 발표", "Streamlit 화면 · TiDB\n질문 35개 · 비교 실험")]
for i, (n, t, d) in enumerate(team):
    x = 0.75 + i * 3.0
    box(s, x, 1.05, 2.8, 1.95, fill=NAVY2, line=None)
    text(s, x + 0.25, 1.2, 2.3, 0.3, f"{n} · [이름]", size=12, color="8A9AAD")
    text(s, x + 0.25, 1.55, 2.3, 0.45, t, size=17, bold=True, color="F2F1EC")
    text(s, x + 0.25, 2.05, 2.3, 0.8, d, size=12, color="B9C3CF", spacing=1.3)
text(s, 0.75, 4.75, 11.8, 1.1, "감사합니다", size=54, bold=True, color="F2F1EC")
text(s, 0.75, 5.95, 11.8, 0.5, "질문 받겠습니다 · github.com/SKNETWORKS-FAMILY-AICAMP/SKN35-3rd-4Team", size=17, color="B9C3CF")
notes(s, "저희 팀은 이렇게 역할을 나눠 진행했습니다. 감사합니다. 질문 받겠습니다. "
         "[예상 질문] 통과하면 정답인가 → 한계 슬라이드, import · 메서드까지 검사 / 35개로 충분한가 → 유형 · 난이도별 경향 확인이 목적 / "
         "왜 LLM에게 검증을 맡기지 않았나 → 검증 LLM도 같은 옛 지식을 가짐 / 근거를 지어낼 수 있지 않나 → verifier가 이번에 보여 준 근거 ID만 인정 / "
         "옛 버전으로 되돌리기는 → 버전별 도구 목록을 비교해서 화면에서 지원")

prs.save(OUT)
print("저장:", OUT, "슬라이드", len(prs.slides))
