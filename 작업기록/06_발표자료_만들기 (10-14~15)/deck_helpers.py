"""발표 PPT 만들기: python build_deck.py <출력.pptx> <화면 캡처.png>"""
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

OUT, SHOT = sys.argv[1], sys.argv[2]

NAVY, NAVY2, LIGHT, WHITE = "14202E", "1E2D3F", "F6F5F1", "FFFFFF"
INK, GRAY, MUTED, LINE = "1B2433", "4A5363", "8A8F98", "E1E0D9"
BLUE, SKY, ORANGE, PEACH = "1F5FAE", "7FB2F0", "B4471B", "F2A27E"
FONT, MONO = "맑은 고딕", "Consolas"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
page = [0]


def rgb(h):
    return RGBColor.from_string(h)


def new(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    page[0] += 1
    return s


def text(s, x, y, w, h, runs, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: 문자열, 또는 줄 목록 [ [(글자, {옵션}), ...], ... ]"""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [[(line, {})] for line in str(runs).split("\n")]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, opt in line:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = rgb(opt.get("color", color))
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", opt.get("font", font))
    return tb


def box(s, x, y, w, h, fill=WHITE, line=LINE, radius=True, dash=False):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                             Inches(x), Inches(y), Inches(w), Inches(h))
    if radius:
        shp.adjustments[0] = 0.06
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1.25 if dash else 0.75)
        if dash:
            shp.line.dash_style = 4
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def arrow(s, x1, y1, x2, y2, color=MUTED, dash=False, width=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def line(s, x1, y1, x2, y2, color=MUTED, dash=False, width=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    return c


def head(s, kicker, title, dark=False, sub=None):
    text(s, 0.75, 0.62, 11.8, 0.3, kicker, size=13, bold=True, color=SKY if dark else BLUE)
    text(s, 0.75, 0.95, 11.8, 0.8, title, size=32, bold=True, color=LIGHT if dark else INK)
    if sub:
        text(s, 0.75, 1.72, 11.8, 0.5, sub, size=16, color="B9C3CF" if dark else GRAY)


def foot(s, dark=False, extra=""):
    text(s, 0.75, 6.92, 11.8, 0.3, f"LangChain 마이그레이션 도우미 · {page[0]}{extra}", size=11,
         color="8A9AAD" if dark else MUTED)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def table(s, x, y, w, rows, col_w, size=15, row_h=0.46, head_fill="ECEBE5", first_bold=False, colors=None):
    shp = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = tblPr.makeelement(qn("a:tableStyleId"), {})
        tblPr.append(style)
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"   # No Style, Table Grid
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(r):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(head_fill if i == 0 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.12)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.text = ""
            run = tf.paragraphs[0].add_run()
            run.text = val
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = i == 0 or (first_bold and j == 0)
            c = INK
            if colors and i > 0 and colors.get(j):
                c = colors[j]
            run.font.color.rgb = rgb(c)
            rPr = run._r.get_or_add_rPr()
            ea = rPr.makeelement(qn("a:ea"), {"typeface": FONT})
            rPr.append(ea)
    return tbl


def big_card(s, x, y, w, h, num, label, src=None, num_color=BLUE, num_size=40):
    box(s, x, y, w, h)
    text(s, x + 0.3, y + 0.25, w - 0.6, 0.8, num, size=num_size, bold=True, color=num_color)
    text(s, x + 0.3, y + 1.1, w - 0.6, 0.9, label, size=15, color=INK, spacing=1.3)
    if src:
        text(s, x + 0.3, y + h - 0.5, w - 0.6, 0.3, src, size=11, color=MUTED)


