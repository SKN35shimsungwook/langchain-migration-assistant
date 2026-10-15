"""발표 대본 Word: python build_docx.py <out.docx>"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from script_v4 import QA, SCRIPT  # noqa: E402

OUT = sys.argv[1]
FONT = "맑은 고딕"
BLUE, GRAY, ORANGE, INK = RGBColor(0x1F, 0x5F, 0xAE), RGBColor(0x6E, 0x6E, 0x73), RGBColor(0xC0, 0x6A, 0x00), RGBColor(0x1B, 0x24, 0x33)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(1.8)

st = doc.styles["Normal"]
st.font.name = FONT
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
st.paragraph_format.space_after = Pt(2)
st.paragraph_format.line_spacing = 1.35


def run(p, t, size=11, bold=False, color=INK):
    r = p.add_run(t)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    return r


def shade(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hex_color)
    tcPr.append(sh)


p = doc.add_paragraph()
run(p, "발표 대본", 24, True)
p = doc.add_paragraph()
run(p, "LangChain 버전 마이그레이션 코드 어시스턴트 · SKN35 3rd 4Team · 15분 (본편 약 14분 20초 + 질문) · 발표 [이름]", 10, color=GRAY)

t = doc.add_table(rows=1, cols=3)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
for cell, (h, d) in zip(t.rows[0].cells, [("화면은 훅만", "숫자 · 한 문장 · 그림을 가리키고, 설명은 말로"),
                                         ("장마다 세 박자", "훅 한 마디 → 설명 → 넘기는 말로 다음 장 예고"),
                                         ("멈추기", "숫자를 말한 뒤, 질문을 던진 뒤 1초 쉬기")]):
    shade(cell, "F2F4F8")
    cp = cell.paragraphs[0]
    run(cp, h + "\n", 10.5, True, BLUE)
    run(cp, d, 9.5, color=GRAY)
doc.add_paragraph()

for i, (tm, scr, hook, body, bridge) in enumerate(SCRIPT, 1):
    tb = doc.add_table(rows=1, cols=1)
    tb.style = "Table Grid"
    trPr = tb.rows[0]._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    cs.set(qn("w:val"), "true")
    trPr.append(cs)
    c = tb.rows[0].cells[0]
    shade(c, "F7F8FA")
    p = c.paragraphs[0]
    run(p, f"{i:02d}  ", 15, True, BLUE)
    run(p, scr, 12.5, True)
    run(p, f"     {tm}", 9.5, color=GRAY)
    p = c.add_paragraph()
    run(p, "훅  ", 11, True, ORANGE)
    run(p, hook, 11.5, True)
    for line in body.split("\n"):
        p = c.add_paragraph()
        run(p, line, 11)
    if bridge:
        p = c.add_paragraph()
        run(p, "→ 넘기며  ", 10.5, True, BLUE)
        run(p, bridge, 10.5, color=BLUE)
    doc.add_paragraph()

doc.add_page_break()
p = doc.add_paragraph()
run(p, "예상 질문과 답", 16, True)
for q, a in QA:
    p = doc.add_paragraph()
    run(p, "Q. " + q, 11, True)
    p = doc.add_paragraph()
    run(p, "A. " + a, 11, color=GRAY)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run(p, "[ ] 칸은 실험 결과 · 이름이 정해지면 채우기", 9, color=GRAY)
doc.save(OUT)
print("저장", OUT)
