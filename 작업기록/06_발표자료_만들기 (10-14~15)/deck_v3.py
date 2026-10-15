"""발표 PPT v3 (애플 플랫, 본편 14장 + 부록 3장): python deck_v3.py <out.pptx> <script.json>"""
import json
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT, SCRIPT_OUT = sys.argv[1], sys.argv[2]
BG, CARD, INK, SUB, FAINT = "F5F5F7", "FFFFFF", "1D1D1F", "6E6E73", "AEAEB2"
BLUE, RED, GREEN, ORANGE, CODE_BG = "0071E3", "FF3B30", "34C759", "FF9F0A", "1D1D1F"
FONT, MONO = "맑은 고딕", "Consolas"
MAIN_TOTAL = 14

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W = 13.333
SCRIPT = []


def rgb(h):
    return RGBColor.from_string(h)


def new(bg=BG):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(bg)
    return s


def text(s, x, y, w, h, t, size=20, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.1):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = t if isinstance(t, list) else [[(ln, {})] for ln in str(t).split("\n")]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for seg, opt in line:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = rgb(opt.get("color", color))
            rPr = r._r.get_or_add_rPr()
            for tag in ("a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is None:
                    el = rPr.makeelement(qn(tag), {})
                    rPr.append(el)
                el.set("typeface", opt.get("font", font))
    return tb


def shadow(shp):
    """부드러운 그림자 (애플 카드 느낌)."""
    spPr = shp._element.spPr
    eff = spPr.makeelement(qn("a:effectLst"), {})
    sh = eff.makeelement(qn("a:outerShdw"), {"blurRad": "254000", "dist": "50800", "dir": "5400000",
                                             "algn": "t", "rotWithShape": "0"})
    clr = sh.makeelement(qn("a:srgbClr"), {"val": "000000"})
    alpha = clr.makeelement(qn("a:alpha"), {"val": "9000"})
    clr.append(alpha)
    sh.append(clr)
    eff.append(sh)
    spPr.append(eff)


def card(s, x, y, w, h, fill=CARD, radius=0.08, soft=True):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.line.fill.background()
    shp.text_frame.text = ""
    if soft:
        shadow(shp)
    else:
        shp.shadow.inherit = False
    return shp


def pill(s, x, y, w, t, fill, color, size=14):
    shp = card(s, x, y, w, 0.46, fill=fill, radius=0.5, soft=False)
    text(s, x, y + 0.08, w, 0.32, t, size=size, bold=True, color=color, align=PP_ALIGN.CENTER)


def arrow(s, x1, y1, x2, y2, color=FAINT, width=2.0, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))


def line(s, x1, y1, x2, y2, color=FAINT, width=2.0, dash=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(width)
    if dash:
        c.line.dash_style = 4


def chrome(s, no, label, appendix=False):
    """작은 머리표: 왼쪽 위 단원, 오른쪽 위 쪽수."""
    text(s, 0.7, 0.5, 6, 0.3, label, size=13, bold=True, color=BLUE if not appendix else SUB)
    text(s, W - 2.7, 0.5, 2.0, 0.3, f"부록 {no}" if appendix else f"{no:02d} / {MAIN_TOTAL}", size=12,
         color=FAINT, align=PP_ALIGN.RIGHT)


def say(s, time, screen, hook, body, bridge):
    """대본: 노트 + 대본 PDF 에 같이 씀."""
    s.notes_slide.notes_text_frame.text = f"[{time}]\n훅: {hook}\n\n{body}\n\n→ 넘기는 말: {bridge}"
    SCRIPT.append({"no": len(prs.slides), "time": time, "screen": screen, "hook": hook, "body": body, "bridge": bridge})


# 01 표지 --------------------------------------------------------------------------
s = new()
text(s, 0.9, 1.0, 11, 0.4, "SKN35 3rd 4Team · LangChain 버전 마이그레이션 코드 어시스턴트", size=15, color=SUB)
text(s, 0.9, 2.3, 11.5, 2.6, "LLM이 준 코드,\n믿어도 될까요?", size=66, bold=True, spacing=1.05)
text(s, 0.9, 5.9, 11, 0.4, "발표 [이름] · 2026-10-15", size=15, color=SUB)
say(s, "0:00 ~ 0:40", "질문 한 줄",
    "ChatGPT가 짜 준 코드, 그대로 실행해 보신 적 있으신가요?",
    "안녕하세요, 4팀 발표를 맡은 [이름]입니다.\n시작하기 전에 하나 여쭤보겠습니다.\nChatGPT가 짜 준 코드, 그대로 실행해 보신 적 있으신가요?\n(잠깐 멈춤)\n잘 돌아가면 다행인데요. 라이브러리 코드는 그렇지 않은 경우가 많습니다.\n오늘은 LLM이 준 코드를 믿어도 되는지에 대한 이야기입니다.",
    "제가 실제로 받아 본 코드를 보여 드리겠습니다.")

# 02 에러 --------------------------------------------------------------------------
s = new()
chrome(s, 2, "문제")
card(s, 1.4, 1.7, 10.5, 3.6, fill=CODE_BG)
text(s, 2.0, 2.2, 9.5, 0.5, ">>> from langchain.chat_models import ChatOpenAI", size=24, color="F5F5F7", font=MONO)
text(s, 2.0, 3.25, 9.5, 1.4, "ImportError:\ncannot import name 'ChatOpenAI'", size=30, bold=True, color="FF6961", font=MONO, spacing=1.15)
text(s, 1.4, 5.7, 10.5, 0.4, "LLM이 준 코드를 LangChain 1.4.4에서 실행한 결과", size=16, color=SUB, align=PP_ALIGN.CENTER)
say(s, "0:40 ~ 1:30", "빨간 에러 한 줄",
    "첫 줄부터 에러입니다.",
    "LLM에게 LangChain 코드를 부탁하면 이런 코드를 줍니다.\n최신 LangChain에서 실행하면, 첫 줄부터 에러입니다.\nLangChain이 1.0에서 패키지를 나누고, 자주 쓰던 LLMChain 같은 기능을 뺐기 때문입니다.\n코드는 맞아 보이는데, 지금은 돌아가지 않는 코드입니다.",
    "이게 저만 겪는 문제일까요?")

# 03 42% ---------------------------------------------------------------------------
s = new()
chrome(s, 3, "문제")
text(s, 0.9, 1.5, 11.5, 2.6, "42%", size=170, bold=True, color=BLUE)
text(s, 0.95, 4.5, 11.5, 0.6, "문서가 부족할 때, LLM 코드가 목표 버전에서 돌아가는 비율", size=26, bold=True)
text(s, 0.95, 5.3, 11.5, 0.4, "When LLMs Lag Behind, 2026 · 문서를 줘도 66%", size=15, color=SUB)
say(s, "1:30 ~ 2:30", "큰 숫자 42%",
    "열 번 중 여섯 번은 실패합니다.",
    "연구 결과를 보면, 문서가 부족할 때 LLM이 만든 코드는 42%만 목표 버전에서 돌아갑니다.\n문서를 줘도 66% 정도입니다.\n이유는 학습 데이터에 있습니다. 인터넷 코드 대부분이 옛 버전이라, LLM은 옛 방식을 답합니다.\n이걸 '버전 지체 환각'이라고 부릅니다.",
    "그런데 진짜 문제는 따로 있습니다.")

# 04 틀려도 모릅니다 ---------------------------------------------------------------------
s = new()
chrome(s, 4, "문제")
text(s, 0.9, 2.4, 11.5, 1.6, "틀려도, 모릅니다.", size=80, bold=True, align=PP_ALIGN.CENTER)
text(s, 0.9, 4.3, 11.5, 0.5, "결국 사람이 에러를 하나씩 검색하며 고칩니다", size=22, color=SUB, align=PP_ALIGN.CENTER)
say(s, "2:30 ~ 3:00", "한 문장",
    "LLM은 자기 답이 틀렸다는 걸 모릅니다.",
    "LLM은 자기 답이 틀렸다는 걸 모릅니다.\n그래서 틀린 코드를 자신 있게 줍니다.\n결국 개발자가 에러를 하나씩 검색하며 고쳐야 합니다.",
    "그래서 저희는 두 가지를 더했습니다.")

# 05 해결 ---------------------------------------------------------------------------
s = new()
chrome(s, 5, "해결")
text(s, 0.9, 1.25, 11.5, 0.9, "그래서, 두 번 확인합니다", size=40, bold=True)
for i, (num, t, d, c) in enumerate([("1", "근거로 고치고", "기억이 아니라\n대응표 · 공식 문서를 보고", BLUE),
                                    ("2", "검사로 확인합니다", "고친 코드를 검사하고\n틀리면 다시 고칩니다", GREEN)]):
    x = 0.9 + i * 6.0
    card(s, x, 2.6, 5.6, 3.5)
    text(s, x + 0.55, 2.95, 1.5, 1.0, num, size=54, bold=True, color=c)
    text(s, x + 0.55, 4.05, 4.6, 0.6, t, size=32, bold=True)
    text(s, x + 0.55, 4.8, 4.6, 1.0, d, size=18, color=SUB, spacing=1.25)
say(s, "3:00 ~ 3:50", "카드 두 장",
    "근거, 그리고 검사. 이 두 가지가 전부입니다.",
    "첫째, 기억이 아니라 근거로 고칩니다. 옛 경로와 새 경로를 잇는 대응표, 그리고 최신 공식 문서를 보고 고칩니다.\n둘째, 고친 코드를 검사합니다. 틀리면 이유를 들고 다시 고칩니다.\n근거, 그리고 검사. 이 두 가지가 저희 챗봇의 전부입니다.",
    "이 두 가지가 어떻게 이어지는지 보여 드리겠습니다.")

# 06 흐름 ---------------------------------------------------------------------------
s = new()
chrome(s, 6, "만든 것")
text(s, 0.9, 1.25, 11.5, 0.9, "네 단계, 실패하면 다시", size=40, bold=True)
steps = [("찾기", "analyzer", "1번"), ("근거", "retriever", "2번"), ("고치기", "rewriter", "3번"), ("검사", "verifier", "3번")]
bw, gap, y0 = 2.5, 0.55, 2.75
x0 = (W - (4 * bw + 3 * gap)) / 2
for i, (t, eng, who) in enumerate(steps):
    x = x0 + i * (bw + gap)
    hl = i == 3
    card(s, x, y0, bw, 1.9, fill=BLUE if hl else CARD)
    text(s, x, y0 + 0.35, bw, 0.7, t, size=32, bold=True, color="FFFFFF" if hl else INK, align=PP_ALIGN.CENTER)
    text(s, x, y0 + 1.15, bw, 0.4, f"{eng} · {who}", size=14, color="DCEBFF" if hl else SUB, align=PP_ALIGN.CENTER)
    if i < 3:
        arrow(s, x + bw + 0.08, y0 + 0.95, x + bw + gap - 0.08, y0 + 0.95)
vx = x0 + 3 * (bw + gap) + bw / 2
rx = x0 + 1 * (bw + gap) + bw / 2
line(s, vx, y0 + 1.9, vx, 5.4, color=ORANGE, width=2.5, dash=True)
line(s, vx, 5.4, rx, 5.4, color=ORANGE, width=2.5, dash=True)
arrow(s, rx, 5.4, rx, y0 + 1.95, color=ORANGE, width=2.5, dash=True)
text(s, rx, 5.55, vx - rx, 0.45, "실패하면 이유를 들고 다시 · 최대 3번", size=18, bold=True, color=ORANGE, align=PP_ALIGN.CENTER)
say(s, "3:50 ~ 5:00", "흐름 그림",
    "검사에 실패하면, 이유를 들고 다시 찾습니다.",
    "흐름은 네 단계입니다.\n먼저 바꿀 곳을 찾고, 근거를 찾고, 그 근거로 고치고, 마지막에 검사합니다.\n검사에 실패하면 실패 이유를 들고 다시 근거를 찾습니다. 이때 틀렸던 근거는 빼고 찾습니다. 최대 세 번입니다.\n이 단계들을 LangGraph로 잇는 일은 1번이 맡았습니다.",
    "그럼 근거는 어디서 올까요?")

# 07 근거 ----------------------------------------------------------------------------
s = new()
chrome(s, 7, "만든 것 · 2번")
text(s, 0.9, 1.5, 11.5, 2.6, "9,155", size=150, bold=True, color=BLUE)
text(s, 0.95, 4.35, 11.5, 0.6, "옛 경로를 새 경로로 잇는 대응표", size=28, bold=True)
for i, (n, t) in enumerate([("295", "공식 문서"), ("11,945", "검색 조각"), ("ID", "바뀐 줄마다 근거")]):
    x = 0.95 + i * 3.3
    text(s, x, 5.35, 3.0, 0.6, n, size=26, bold=True)
    text(s, x, 5.95, 3.0, 0.4, t, size=15, color=SUB)
say(s, "5:00 ~ 6:00", "큰 숫자 9,155",
    "옛 경로와 새 경로를 잇는 대응표가 9,155줄입니다.",
    "근거는 2번이 만들었습니다.\n옛 import 경로를 새 경로로 잇는 대응표가 9,155줄입니다.\n여기에 공식 문서 295개를 잘게 나눠, 필요한 부분을 바로 찾을 수 있게 했습니다.\n그리고 바뀐 줄마다 어떤 근거를 썼는지 ID가 붙습니다. 나중에 왜 이렇게 바꿨는지 따라갈 수 있습니다.",
    "근거로 고친 코드, 이제 검사합니다.")

# 08 검사 ----------------------------------------------------------------------------
s = new()
chrome(s, 8, "만든 것 · 3번")
text(s, 0.9, 1.5, 4.0, 2.6, "11", size=170, bold=True, color=GREEN)
text(s, 4.2, 2.2, 8.0, 0.8, "가지 검사", size=44, bold=True)
text(s, 4.25, 3.1, 8.0, 0.6, "LLM 없이, 코드로", size=30, color=SUB)
for i, t in enumerate(["새 경로가 진짜 있나", "기록 없이 바뀐 줄은 없나", "결과 모양이 그대로인가"]):
    pill(s, 0.95 + i * 3.85, 5.2, 3.6, t, fill="E8F7EC", color="1E7B34", size=15)
text(s, 0.95, 5.9, 11, 0.4, "외 8가지 · 부록 2", size=13, color=FAINT)
say(s, "6:00 ~ 7:00", "큰 숫자 11",
    "검사에는 LLM을 쓰지 않습니다.",
    "검사는 3번이 만들었습니다. 모두 11가지인데, LLM을 쓰지 않습니다.\n이유는 간단합니다. 검사하는 LLM도 같은 옛 지식을 갖고 있기 때문입니다.\n그래서 코드로 확인합니다. 새 경로가 정말 있는지, 근거 없이 바뀐 줄은 없는지, 결과 모양이 그대로인지 봅니다.\n근거 없이 바뀐 줄은 곧 환각이라서, 바로 걸러 냅니다.",
    "이 결과를 사용자는 어떻게 볼까요?")

# 09 화면 ----------------------------------------------------------------------------
s = new()
chrome(s, 9, "만든 것 · 4번")
text(s, 0.9, 1.15, 11.5, 0.8, "결과와 근거를 한 화면에", size=36, bold=True)
c = card(s, 0.9, 2.15, 11.5, 4.65)
line(s, 0.9, 2.15, 0.9, 2.15)
text(s, 0.9, 4.0, 11.5, 0.9, "[화면 캡처]\n확인용 예제를 보낸 뒤 결과 패널 (검증 배지 · 전후 비교 빨강 · 초록)", size=18, bold=True, color=FAINT,
     align=PP_ALIGN.CENTER)
say(s, "7:00 ~ 8:00", "화면 캡처",
    "코드를 넣고, 말로 부탁하면 됩니다.",
    "제가 맡은 화면입니다.\n코드를 넣고 '0.2로 되돌려 줘'처럼 말로 부탁하면 됩니다. 버전을 적지 않으면 최신으로 바꿉니다.\n기다리는 동안 지금 어느 단계인지, 검사에서 떨어져 다시 고치는 중인지가 실시간으로 보입니다.\n결과가 나오면 지운 줄은 빨강, 새 줄은 초록으로 보이고, 근거와 11가지 검사 결과도 탭으로 볼 수 있습니다.\n대화는 TiDB에 저장돼서, 언제든 다시 열어 볼 수 있습니다.",
    "실제로 돌려 본 결과를 보여 드리겠습니다.")

# 10 실제 사례 ------------------------------------------------------------------------
s = new()
chrome(s, 10, "결과 · 실제 실행")
text(s, 0.9, 1.15, 8.5, 0.8, "세 번째에 통과했습니다", size=36, bold=True)
pill(s, 9.6, 1.3, 2.8, "검사 통과 · 3회", fill="E8F7EC", color="1E7B34", size=15)
for i, (label, code, col) in enumerate([
        ("이전", "from langchain.chains import LLMChain\n\nchain = LLMChain(llm=llm, prompt=prompt)\nchain.run(word='크다')", "FF6961"),
        ("이후", "from langchain_core.output_parsers \\\n    import StrOutputParser\nchain = prompt | llm | StrOutputParser()\nchain.invoke({'word': '크다'})", "64D2FF")]):
    x = 0.9 + i * 5.9
    card(s, x, 2.3, 5.6, 3.2, fill=CODE_BG)
    text(s, x + 0.4, 2.55, 4.8, 0.4, label, size=16, bold=True, color=col)
    text(s, x + 0.4, 3.15, 5.0, 2.2, code, size=15, color="F5F5F7", font=MONO, spacing=1.3)
text(s, 0.9, 5.85, 11.5, 0.5, "바뀐 6줄 모두 근거 ID  ·  못 바꾼 줄 0  ·  약 1분", size=19, color=SUB, align=PP_ALIGN.CENTER)
say(s, "8:00 ~ 9:10", "이전 · 이후 코드",
    "처음 두 번은 떨어지고, 세 번째에 통과했습니다.",
    "어제 실제로 돌려 본 결과입니다.\n왼쪽은 예전 방식의 LLMChain 코드입니다. 오른쪽처럼 최신 방식으로 바뀌었습니다.\n중요한 건 과정입니다. 처음 두 번은 검사에 걸렸고, 실패 이유를 들고 다시 고쳐서 세 번째에 통과했습니다.\n바뀐 여섯 줄 모두 근거가 붙었습니다. 같은 코드를 두 번 돌렸고, 두 번 모두 세 번째에 통과했습니다.",
    "그럼 일반 LLM보다 얼마나 나을까요?")

# 11 평가 설계 ------------------------------------------------------------------------
s = new()
chrome(s, 11, "결과 · 평가")
text(s, 0.9, 1.5, 4.5, 2.6, "30", size=170, bold=True, color=BLUE)
text(s, 4.6, 2.2, 8.0, 0.8, "개의 실무형 질문", size=44, bold=True)
text(s, 4.65, 3.1, 8.0, 0.6, "실제 사용자처럼, 요청 문장 + 옛 코드", size=24, color=SUB)
for i, t in enumerate(["같은 질문", "같은 모델", "같은 채점"]):
    pill(s, 0.95 + i * 3.0, 5.2, 2.7, t, fill="E5F0FF", color=BLUE, size=17)
say(s, "9:10 ~ 10:00", "큰 숫자 30",
    "같은 질문, 같은 모델, 같은 채점으로 비교했습니다.",
    "일반 LLM과 비교하려고, 실제 사용자가 보낼 법한 질문 30개를 만들었습니다.\n요청 문장과 옛 코드가 함께 들어 있고, 열 가지 유형을 세 개씩 담았습니다.\n30개 모두 최신 버전에서 실제로 깨지는 코드인지 미리 확인했습니다.\n비교는 공정하게 했습니다. 같은 질문, 같은 모델, 같은 채점 기준입니다.",
    "결과입니다.")

# 12 결과 ----------------------------------------------------------------------------
s = new()
chrome(s, 12, "결과")
text(s, 0.9, 1.15, 11.5, 0.8, "검사 통과율", size=36, bold=True)
for i, (who, col) in enumerate([("우리 챗봇", BLUE), ("일반 LLM", FAINT)]):
    x = 0.9 + i * 6.0
    card(s, x, 2.3, 5.6, 3.6)
    text(s, x + 0.5, 2.7, 4.6, 0.5, who, size=22, bold=True, color=SUB)
    text(s, x + 0.5, 3.3, 4.8, 2.0, "[__]%", size=96, bold=True, color=col)
text(s, 0.9, 6.15, 11.5, 0.4, "질문 30개 · 같은 모델 · verifier와 같은 채점 기준", size=15, color=FAINT, align=PP_ALIGN.CENTER)
say(s, "10:00 ~ 11:00", "숫자 두 개",
    "[결과 한 문장: 예) 일반 LLM의 [__]배입니다]",
    "[실험 결과가 나오면 채워 주세요]\n일반 LLM은 [__]%, 저희 챗봇은 [__]%였습니다.\n특히 [__] 유형에서 차이가 가장 컸습니다.\n다시 고치는 과정 덕분에, 첫 시도 [__]%에서 [__]%까지 올라갔습니다.",
    "물론 한계도 있습니다.")

# 13 한계 ----------------------------------------------------------------------------
s = new()
chrome(s, 13, "한계")
text(s, 0.9, 2.2, 11.5, 1.6, "통과 ≠ 정답", size=88, bold=True, align=PP_ALIGN.CENTER)
for i, (t, d) in enumerate([("다음 1", "실제 실행 결과까지 확인"), ("다음 2", "1분 걸리는 변환을 더 빠르게")]):
    x = 2.2 + i * 4.6
    card(s, x, 4.45, 4.3, 1.35)
    text(s, x + 0.4, 4.62, 3.6, 0.4, t, size=14, bold=True, color=BLUE)
    text(s, x + 0.4, 5.05, 3.7, 0.6, d, size=20, bold=True)
say(s, "11:00 ~ 12:10", "한 문장 + 다음 단계 둘",
    "검사를 통과했다고, 정답은 아닙니다.",
    "한계도 분명히 말씀드리겠습니다.\n검사를 통과했다고 정답은 아닙니다. 지금은 코드를 검사할 뿐, 실제로 실행한 결과가 맞는지는 아직 보지 않습니다.\n또 검사하고 다시 고치느라, 한 번 바꾸는 데 1분 정도 걸립니다.\n다음 단계는 실제 실행까지 확인하는 것, 그리고 변환을 더 빠르게 만드는 것입니다.",
    "마지막으로 정리하겠습니다.")

# 14 마무리 ---------------------------------------------------------------------------
s = new()
text(s, 0.9, 1.4, 11.5, 2.4, "라이브러리가 바뀌어도,\n근거와 함께.", size=60, bold=True, spacing=1.08)
text(s, 0.95, 4.0, 11, 0.5, "감사합니다 · 질문 받겠습니다", size=24, color=SUB)
for i, (n, t) in enumerate([("1번", "그래프"), ("2번", "데이터 · 검색"), ("3번", "변환 · 검증"), ("4번", "화면 · DB · 평가")]):
    x = 0.95 + i * 2.95
    card(s, x, 5.05, 2.7, 1.05)
    text(s, x + 0.3, 5.2, 2.2, 0.35, f"{n} · [이름]", size=12, color=FAINT)
    text(s, x + 0.3, 5.55, 2.3, 0.4, t, size=17, bold=True)
text(s, 0.95, 6.5, 11.5, 0.35, "github.com/SKN35shimsungwook/langchain-migration-assistant  ·  [배포 주소]", size=12, color=FAINT)
say(s, "12:10 ~ 13:00 (+ 질문)", "마무리 한 문장",
    "라이브러리가 바뀌어도, 근거와 함께.",
    "라이브러리는 앞으로도 계속 바뀝니다.\n저희 챗봇은 그때마다, 근거와 함께 검사한 코드를 드립니다.\n1번부터 4번까지 네 명이 그래프, 데이터, 검증, 화면을 나눠 만들었습니다.\n감사합니다. 질문 받겠습니다.",
    "(질문이 나오면 부록 1~3장을 꺼내 답하기)")

# 부록 1: ERD ------------------------------------------------------------------------
s = new()
chrome(s, 1, "부록 · 대화 기록 (TiDB)", appendix=True)
text(s, 0.9, 1.15, 11.5, 0.8, "사용자 한 명, 대화 여러 개", size=32, bold=True)
def entity(x, y, name, rows, w=5.2):
    card(s, x, y, w, 0.75 + 0.5 * len(rows))
    text(s, x + 0.4, y + 0.22, w - 0.8, 0.4, name, size=20, bold=True, color=BLUE)
    for i, (col, typ, mark) in enumerate(rows):
        yy = y + 0.8 + i * 0.5
        text(s, x + 0.4, yy, 2.0, 0.4, col, size=15, bold=bool(mark), font=MONO)
        text(s, x + 2.4, yy, 1.9, 0.4, typ, size=12, color=SUB, font=MONO)
        text(s, x + 4.3, yy, 0.8, 0.4, mark, size=12, bold=True, color=ORANGE)
entity(0.9, 2.3, "users", [("uuid", "VARCHAR(36)", "PK"), ("user_id", "VARCHAR(100)", "UQ"), ("created_at", "DATETIME", "")])
entity(7.2, 2.3, "chat_history", [("uuid", "VARCHAR(36)", "PK"), ("user_id", "VARCHAR(100)", "FK"),
                                   ("session_id", "VARCHAR(100)", "IDX"), ("question", "TEXT", ""),
                                   ("answer", "TEXT", ""), ("created_at", "DATETIME", "")])
arrow(s, 6.1, 3.6, 7.2, 3.6, color=BLUE, width=2.5)
text(s, 6.1, 3.15, 1.1, 0.4, "1 : N", size=14, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
text(s, 0.9, 5.0, 5.2, 1.3, "질문 + 답 = 한 행\n새 대화 = 새 session_id\n가입 때 아이디 중복 검사", size=16, color=SUB, spacing=1.4)
say(s, "질문용", "ERD", "대화 기록은 이렇게 저장합니다.",
    "users에는 아이디가 중복 없이 들어가고, chat_history에는 질문 하나와 답 하나가 한 행으로 들어갑니다.\n새 대화를 누르면 새 session_id가 생기고, 목록은 session_id로 묶어 보여 줍니다.\nTiDB에 연결할 수 없으면 화면이 멈추지 않고 로컬 저장으로 넘어갑니다.", "")

# 부록 2: 검사 11가지 -------------------------------------------------------------------
s = new()
chrome(s, 2, "부록 · 검사 11가지 (3번)", appendix=True)
text(s, 0.9, 1.15, 11.5, 0.8, "LLM 없이, 코드로 11가지", size=32, bold=True)
checks = ["문법", "근거 ID", "변경 기록", "남은 옛 import", "import 존재", "남은 옛 메서드",
          "안 쓰는 import", "반환값 보존", "정의 안 된 이름", "출력 꺼내기", "체인 입출력 키"]
for i, t in enumerate(checks):
    col, row = i % 4, i // 4
    x, y = 0.9 + col * 2.9, 2.2 + row * 1.35
    card(s, x, y, 2.65, 1.1)
    text(s, x + 0.3, y + 0.17, 0.6, 0.4, f"{i + 1}", size=16, bold=True, color=GREEN)
    text(s, x + 0.3, y + 0.55, 2.3, 0.4, t, size=17, bold=True)
say(s, "질문용", "검사 11가지", "검사 항목 전체입니다.",
    "문법, 근거 ID가 이번에 보여 준 근거인지, 기록 없이 바뀐 줄이 없는지, 옛 import가 남았는지, 새 경로가 1.x에 있는지,\n옛 메서드, 안 쓰는 import, 반환값 보존, 정의 안 된 이름, 출력 꺼내기, 체인 입출력 키까지 봅니다.\n오류는 다시 고칠 이유가 되고, 경고는 결과에 표시만 합니다.", "")

# 부록 3: 질문 구성 ----------------------------------------------------------------------
s = new()
chrome(s, 3, "부록 · 평가 질문", appendix=True)
text(s, 0.9, 1.15, 11.5, 0.8, "10가지 유형 × 3개 = 30개", size=32, bold=True)
types = ["import 경로", "로더 · 분할기", "임베딩", "LLMChain → LCEL", "체인 조합", "RAG", "메모리", "에이전트", "출력 파서", "실무 상황"]
for i, t in enumerate(types):
    col, row = i % 5, i // 5
    x, y = 0.9 + col * 2.32, 2.25 + row * 1.0
    card(s, x, y, 2.12, 0.78)
    text(s, x, y, 2.12, 0.78, t, size=15, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
for i, (n, t) in enumerate([("7", "쉬움"), ("15", "중간"), ("8", "어려움"), ("36/36", "미리 확인 · 버전 오가기 6 포함")]):
    x = 0.95 + i * 2.6
    text(s, x, 4.75, 4.0 if i == 3 else 2.5, 0.7, n, size=34, bold=True, color=BLUE)
    text(s, x, 5.45, 4.0 if i == 3 else 2.5, 0.6, t, size=14, color=SUB)
say(s, "질문용", "질문 구성", "평가 질문 구성입니다.",
    "열 가지 유형을 세 개씩, 쉬움 7, 중간 15, 어려움 8입니다.\n옛 버전으로 되돌리는 질문 6개까지 36개 모두, 옛 코드가 실제로 깨지는지와 기대하는 답이 실제로 되는지 미리 확인했습니다.", "")

prs.save(OUT)
json.dump(SCRIPT, open(SCRIPT_OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("저장", OUT, len(prs.slides), "장")
